#!/usr/bin/env python3
"""Export every booked leg and the P&L that follows from them to one .xlsx.

Built from the local journal and ledger rather than the Google Sheet, so the
file is current even when a sheet write has not run yet. Open positions are
marked at their stored futures expiry's last traded price when Angel One is
reachable; otherwise that tab reports entry prices only.

    .venv/bin/python scripts/export_pnl_workbook.py
"""

from __future__ import annotations

import csv
import json
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.config import load_config
from src.paper_trading.journal import DATE_FORMAT

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=12)
GREEN = Font(color="107C41")
RED = Font(color="C00000")
BOLD = Font(bold=True)
RULE = Border(top=Side(style="thin", color="BFBFBF"))
MONEY = "#,##0"

REASON_LABEL = {
    "first_target": "Lot 1 booked — 5% or SMMA 21",
    "second_target": "Lot 2 booked — 12% or SMMA 50",
    "third_target": "Third target",
    "rsi_target": "Final lot — RSI 30 / 70",
    "smma_cross": "Final lot — strong candle back through SMMA 21",
    "stop_loss": "Stopped out",
    "expiry": "Contract expiry",
    "wall_broken": "OI wall broken",
    "strike_through": "Cash through entry strike",
    "writing_gone": "Option writing unwound",
}


def money(cell, value) -> None:
    cell.value = value
    cell.number_format = MONEY
    if isinstance(value, (int, float)):
        cell.font = GREEN if value >= 0 else RED


def leg_charges(entry: float, exit_: float, margin: float, side: str,
                margin_pct: float) -> float:
    """Round-trip cost of one leg under the current NSE stock-futures schedule."""
    if entry <= 0 or margin <= 0:
        return 0.0
    entry_notional = margin / (margin_pct / 100)
    exit_notional = entry_notional * (exit_ / entry)
    if side == "Sell":
        sell, buy = entry_notional, exit_notional
    else:
        buy, sell = entry_notional, exit_notional
    turnover = buy + sell
    brokerage = 40.0
    stt = sell * 0.0002
    stamp = buy * 0.00002
    exchange = turnover * 0.0000188
    sebi = turnover * 10 / 10_000_000
    clearing = turnover * 0.000005
    gst = (brokerage + exchange + sebi + clearing) * 0.18
    return brokerage + stt + stamp + exchange + sebi + clearing + gst


def read_legs(path: Path, margin_pct: float) -> list[dict]:
    legs: list[dict] = []
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            entry = float(row["Entry price"] or 0)
            exit_ = float(row["Exit price"] or 0)
            per_lot, _, lots = str(row["Capital needed"]).partition("*")
            margin = float(per_lot or 0) * (float(lots) if lots else 1.0)
            row["_pnl"] = float(row["Profit/loss"] or 0)
            row["_margin"] = margin
            row["_lots"] = int(float(lots)) if lots else 1
            row["_charges"] = leg_charges(
                entry, exit_, margin, row["Buy/Sell"], margin_pct
            )
            row["_return"] = (
                (exit_ - entry) / entry * (1 if row["Buy/Sell"] == "Buy" else -1) * 100
                if entry else 0.0
            )
            row["_exit_dt"] = datetime.strptime(row["Exit date"], DATE_FORMAT)
            row["_entry_dt"] = datetime.strptime(row["Entry date"], DATE_FORMAT)
            legs.append(row)
    legs.sort(key=lambda r: (r["_exit_dt"], r["_entry_dt"], r["Symbol"]))
    return legs


def write_table(sheet, headers, rows, *, money_cols=(), pct_cols=(), start=1):
    for c, name in enumerate(headers, start=1):
        cell = sheet.cell(row=start, column=c, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start=start + 1):
        for c, value in enumerate(row, start=1):
            cell = sheet.cell(row=r, column=c)
            if c - 1 in money_cols:
                money(cell, value)
            elif c - 1 in pct_cols:
                cell.value = value
                cell.number_format = '0.0"%"'
            else:
                cell.value = value
    widths: dict[int, int] = {}
    for row in [headers, *rows]:
        for c, value in enumerate(row, start=1):
            widths[c] = min(max(widths.get(c, 10), len(str(value)) + 2), 62)
    for c, width in widths.items():
        sheet.column_dimensions[get_column_letter(c)].width = width
    sheet.freeze_panes = sheet.cell(row=start + 1, column=1).coordinate


def open_position_marks(positions: list[dict]) -> dict[str, float]:
    try:
        from src.data.angelone_client import AngelOneClient
    except Exception:
        return {}
    config = load_config(ROOT / "config.yaml")
    client = AngelOneClient(
        rsi_period=config.rsi.period,
        history_days=config.data.history_days,
        extreme_history_days=config.data.extreme_history_days,
    )
    try:
        pairs = [(p["symbol"], p.get("expiry", "")) for p in positions]
        return client.get_futures_ltps_for_expiries(pairs)
    except Exception as exc:
        print(f"  could not price open positions ({exc}) — showing entries only")
        return {}
    finally:
        client.close()


def main() -> None:
    config = load_config(ROOT / "config.yaml")
    paper = config.rsi_candle_2w_paper_trading
    legs = read_legs(Path(paper.journal_csv), paper.margin_pct)
    ledger = json.loads(Path(paper.ledger_path).read_text())
    positions = [p for p in ledger["positions"] if p.get("lots_open", 0) > 0]

    journal_total = sum(leg["_pnl"] for leg in legs)
    ledger_total = float(ledger["realised_pnl"])
    unrecorded = ledger_total - journal_total
    charges = sum(leg["_charges"] for leg in legs)

    print(f"  {len(legs)} booked legs, {len(positions)} positions still open")
    marks = open_position_marks(positions)
    print(f"  priced {len(marks)} of {len(positions)} open positions")

    workbook = Workbook()

    # ---------------------------------------------------------------- Summary
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "RSI_CandlePattern — realised P&L"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"Generated {datetime.now():%d %b %Y %H:%M}"
    sheet["A2"].font = Font(color="808080")

    wins = [leg for leg in legs if leg["_pnl"] > 0]
    losses = [leg for leg in legs if leg["_pnl"] <= 0]
    gross_win = sum(leg["_pnl"] for leg in wins)
    gross_loss = sum(leg["_pnl"] for leg in losses)
    unreal = 0.0
    for position in positions:
        mark = marks.get(position["symbol"])
        if not mark:
            continue
        move = mark - position["entry_price"]
        if position["direction"] == "SHORT":
            move = -move
        unreal += move * position["lot_size"] * position["lots_open"]

    lines = [
        ("Booked legs in the journal", len(legs), None),
        ("Winning legs", len(wins), None),
        ("Losing legs", len(losses), None),
        ("Win rate", f"{len(wins) / len(legs) * 100:.0f}%" if legs else "—", None),
        ("", "", None),
        ("Gross profit on winning legs", gross_win, "money"),
        ("Gross loss on losing legs", gross_loss, "money"),
        ("Net booked P&L in this journal", journal_total, "money"),
        ("Legs missing from the journal (see note)", unrecorded, "money"),
        ("Realised P&L per the ledger", ledger_total, "money"),
        ("", "", None),
        ("Estimated charges on booked legs", -charges, "money"),
        ("Realised P&L after charges", ledger_total - charges, "money"),
        ("", "", None),
        (f"Open positions ({len(positions)})", "", None),
        ("Unrealised P&L at last traded price", unreal, "money"),
        ("Margin blocked",
         sum(float(p.get("margin_blocked") or 0) for p in positions), "money"),
        ("", "", None),
        ("Total including open positions", ledger_total - charges + unreal, "money"),
    ]
    for r, (label, value, kind) in enumerate(lines, start=4):
        sheet.cell(row=r, column=1, value=label)
        cell = sheet.cell(row=r, column=2)
        if kind == "money":
            money(cell, value)
        else:
            cell.value = value
        if label.startswith(("Realised P&L per", "Total including")):
            sheet.cell(row=r, column=1).font = BOLD
            cell.font = BOLD
            sheet.cell(row=r, column=1).border = RULE
            cell.border = RULE

    note = sheet.cell(
        row=len(lines) + 6,
        column=1,
        value=(
            "Note: on 3 Sep the journal CSV was truncated and 41 already-booked legs "
            "lost their rows. Their P&L is still counted in the ledger total, which is "
            "why the two figures differ by the amount shown above. Those individual "
            "rows cannot be recovered."
        ),
    )
    note.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(
        start_row=note.row, start_column=1, end_row=note.row + 2, end_column=2
    )
    sheet.column_dimensions["A"].width = 46
    sheet.column_dimensions["B"].width = 22

    # --------------------------------------------------------- Booked trades
    sheet = workbook.create_sheet("Booked trades")
    write_table(
        sheet,
        ["Symbol", "Buy/Sell", "Entry date", "Entry price", "Entry RSI", "Exit date",
         "Exit price", "Exit RSI", "Days held", "Lots", "Margin", "Return %",
         "Profit/loss", "Charges", "Net", "Why it was booked"],
        [[
            leg["Symbol"], leg["Buy/Sell"], leg["Entry date"], float(leg["Entry price"]),
            leg["Entry RSI"], leg["Exit date"], float(leg["Exit price"]), leg["Exit RSI"],
            leg["Holding trading period"], leg["_lots"], round(leg["_margin"]),
            round(leg["_return"], 2), leg["_pnl"], -round(leg["_charges"]),
            round(leg["_pnl"] - leg["_charges"]),
            leg["Exit trigger"] or REASON_LABEL.get(leg["Exit reason"], leg["Exit reason"]),
        ] for leg in legs],
        money_cols=(10, 12, 13, 14),
        pct_cols=(11,),
    )
    sheet.auto_filter.ref = sheet.dimensions

    # ----------------------------------------------------------- By exit rule
    by_reason: dict[str, list[dict]] = defaultdict(list)
    for leg in legs:
        by_reason[leg["Exit reason"]].append(leg)
    sheet = workbook.create_sheet("By exit rule")
    write_table(
        sheet,
        ["Exit rule", "Legs", "Winners", "Total P&L", "Average per leg"],
        [[
            REASON_LABEL.get(reason, reason), len(rows),
            sum(1 for r in rows if r["_pnl"] > 0),
            round(sum(r["_pnl"] for r in rows)),
            round(sum(r["_pnl"] for r in rows) / len(rows)),
        ] for reason, rows in sorted(
            by_reason.items(), key=lambda kv: sum(r["_pnl"] for r in kv[1]), reverse=True
        )],
        money_cols=(3, 4),
    )

    # --------------------------------------------------------------- By month
    by_month: dict[str, list[dict]] = defaultdict(list)
    for leg in legs:
        by_month[f"{leg['_exit_dt']:%Y-%m}"].append(leg)
    running = 0.0
    month_rows = []
    for month in sorted(by_month):
        rows = by_month[month]
        total = sum(r["_pnl"] for r in rows)
        running += total
        month_rows.append([
            f"{datetime.strptime(month, '%Y-%m'):%b %Y}", len(rows),
            sum(1 for r in rows if r["_pnl"] > 0), round(total), round(running),
        ])
    sheet = workbook.create_sheet("By month")
    write_table(
        sheet,
        ["Month booked", "Legs", "Winners", "P&L", "Cumulative"],
        month_rows,
        money_cols=(3, 4),
    )

    # -------------------------------------------------------------- By symbol
    by_symbol: dict[str, list[dict]] = defaultdict(list)
    for leg in legs:
        by_symbol[leg["Symbol"]].append(leg)
    sheet = workbook.create_sheet("By symbol")
    write_table(
        sheet,
        ["Symbol", "Legs", "Winners", "Total P&L", "Best leg", "Worst leg"],
        [[
            symbol, len(rows), sum(1 for r in rows if r["_pnl"] > 0),
            round(sum(r["_pnl"] for r in rows)),
            round(max(r["_pnl"] for r in rows)),
            round(min(r["_pnl"] for r in rows)),
        ] for symbol, rows in sorted(
            by_symbol.items(), key=lambda kv: sum(r["_pnl"] for r in kv[1]), reverse=True
        )],
        money_cols=(3, 4, 5),
    )

    # -------------------------------------------------------- Open positions
    open_rows = []
    for position in sorted(positions, key=lambda p: p["entry_time"]):
        mark = marks.get(position["symbol"])
        pnl = ""
        if mark:
            move = mark - position["entry_price"]
            if position["direction"] == "SHORT":
                move = -move
            pnl = round(move * position["lot_size"] * position["lots_open"])
        booked = sum(leg["pnl"] for leg in position.get("closed_legs", []))
        open_rows.append([
            position["symbol"],
            "Sell" if position["direction"] == "SHORT" else "Buy",
            position["entry_time"][:10],
            position["entry_price"],
            f"{position['lots_open']} of {position['lots_total']}",
            position.get("expiry", ""),
            mark or "",
            pnl,
            round(booked) if booked else "",
            round(float(position.get("margin_blocked") or 0)),
            position.get("stop_price") or "no stop",
        ])
    sheet = workbook.create_sheet("Open positions")
    write_table(
        sheet,
        ["Symbol", "Buy/Sell", "Entry date", "Entry price", "Lots open", "Expiry",
         "Last price", "Unrealised", "Already booked", "Margin", "Stop"],
        open_rows,
        money_cols=(7, 8, 9),
    )
    sheet.auto_filter.ref = sheet.dimensions

    out = ROOT / "data" / f"RSI_CandlePattern_PnL_{datetime.now():%Y-%m-%d}.xlsx"
    workbook.save(out)
    print(f"\nWrote {out}  ({out.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
