#!/usr/bin/env python3
"""Export the RSI_CandlePattern backtest results to one .xlsx workbook.

Reads whatever result files scripts/backtest_rsi_candle_1y.py has written and
lays each run out as its own tab: the headline comparison, realised P&L by
month, and the legs each exit rule closed.

    .venv/bin/python scripts/export_backtest_workbook.py
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CACHE = ROOT / "data" / "backtests" / "rsi_candle_1y"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=12)
GREY = Font(color="808080")
GREEN = Font(color="107C41")
RED = Font(color="C00000")
BOLD = Font(bold=True)
HIGHLIGHT = PatternFill("solid", fgColor="E2EFDA")
RULE = Border(top=Side(style="thin", color="BFBFBF"))
MONEY = "#,##0"

RUNS = [
    ("Full year", "results_ladder.json",
     "Three-lot ladder vs the two-lot book, 8 Sep 2025 – 11 Sep 2026."),
    ("From 1 May", "results_ladder_2026-05-01.json",
     "The same comparison over 4 May – 11 Sep 2026 only."),
    ("Stop levels", "results.json",
     "Entry-candle stop vs flat percent stops, on the two-lot book."),
]

REASON_LABEL = {
    "first_target": "Lot 1 — 5% or SMMA 21",
    "second_target": "Lot 2 — 12% or SMMA 50",
    "rsi_target": "Final lot — RSI 30 / 70",
    "smma_cross": "Final lot — SMMA 21 cross",
    "stop_loss": "Stopped out",
    "expiry": "Contract expiry",
}

# Rows worth drawing the eye to on each tab.
FEATURED = {"3 lot 4Cr ladder", "pct-2"}


def money(cell, value):
    cell.value = value
    cell.number_format = MONEY
    if isinstance(value, (int, float)):
        cell.font = GREEN if value >= 0 else RED
    return cell


def autosize(sheet, rows):
    widths: dict[int, int] = {}
    for row in rows:
        for c, value in enumerate(row, start=1):
            widths[c] = min(max(widths.get(c, 10), len(str(value)) + 2), 60)
    for c, width in widths.items():
        sheet.column_dimensions[get_column_letter(c)].width = width


def header(sheet, row, names):
    for c, name in enumerate(names, start=1):
        cell = sheet.cell(row=row, column=c, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_run(workbook, title, payload, blurb):
    results = payload["results"]
    start, end = payload["window"]

    sheet = workbook.create_sheet(title[:31])
    sheet["A1"] = f"{title} — {start} to {end}"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        f"{blurb}  {payload['names']} F&O names, {payload['signals']:,} signals. "
        "Net is after brokerage, STT, stamp duty, exchange and SEBI fees, "
        "clearing and GST."
    )
    sheet["A2"].font = GREY
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)

    cols = ["Book", "Capital", "Lots", "Trades", "Legs", "% killed by stop",
            "Leg win rate", "Realised", "Unrealised", "Open at end", "Charges",
            "Net P&L", "Return on capital", "Worst drawdown", "Peak margin used"]
    header(sheet, 4, cols)

    table = [cols]
    for r, item in enumerate(results, start=5):
        legs = item["legs"] or 1
        # The stop-level run predates the 3-lot work and stored neither field.
        capital = item.get("capital", 20_000_000)
        row = [
            item["variant"],
            f"₹{capital / 10_000_000:.0f} Cr",
            item.get("lots", 2),
            item["opened"],
            item["legs"],
            item["stops"] / legs,
            item["win_legs"] / legs,
            item["realised"],
            item["unrealised"],
            item["still_open"],
            -item["charges"],
            item["net_total"],
            item["net_total"] / capital,
            item["max_drawdown"],
            item["peak_margin"],
        ]
        table.append(row)
        for c, value in enumerate(row, start=1):
            cell = sheet.cell(row=r, column=c)
            if c in (6, 7, 13):
                cell.value = value
                cell.number_format = "0.0%"
            elif c in (8, 9, 11, 12, 14, 15):
                money(cell, value)
            else:
                cell.value = value
            if item["variant"] in FEATURED:
                cell.fill = HIGHLIGHT
                if c in (1, 12):
                    cell.font = BOLD if c == 1 else Font(bold=True, color="107C41")
    autosize(sheet, table)
    sheet.freeze_panes = "B5"

    # ------------------------------------------------- realised P&L by month
    months = sorted({m for item in results for m in item["monthly"]})
    top = len(results) + 7
    sheet.cell(row=top, column=1, value="Realised P&L by month booked").font = BOLD
    sheet.cell(row=top, column=1).border = RULE
    header(sheet, top + 1, ["Book", *[
        f"{datetime.strptime(m, '%Y-%m'):%b %y}" for m in months
    ], "Total"])
    for r, item in enumerate(results, start=top + 2):
        sheet.cell(row=r, column=1, value=item["variant"])
        for c, month in enumerate(months, start=2):
            money(sheet.cell(row=r, column=c), item["monthly"].get(month, 0))
        money(sheet.cell(row=r, column=len(months) + 2),
              sum(item["monthly"].values())).font = BOLD

    # -------------------------------------------------------- legs by reason
    reasons = sorted({k for item in results for k in item.get("by_reason", {})})
    top = top + len(results) + 5
    sheet.cell(row=top, column=1,
               value="Legs closed by each exit rule (count, then P&L)").font = BOLD
    sheet.cell(row=top, column=1).border = RULE
    names = [REASON_LABEL.get(k, k) for k in reasons]
    header(sheet, top + 1, ["Book", *[f"{n}\nlegs" for n in names],
                            *[f"{n}\nP&L" for n in names]])
    for r, item in enumerate(results, start=top + 2):
        sheet.cell(row=r, column=1, value=item["variant"])
        for c, key in enumerate(reasons, start=2):
            sheet.cell(row=r, column=c,
                       value=item.get("by_reason", {}).get(key, [0, 0])[0])
        for c, key in enumerate(reasons, start=2 + len(reasons)):
            money(sheet.cell(row=r, column=c),
                  item.get("by_reason", {}).get(key, [0, 0])[1])

    return len(results)


def main() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    notes = workbook.create_sheet("Read me")
    lines = [
        ("RSI_CandlePattern backtest results", TITLE_FONT),
        (f"Generated {datetime.now():%d %b %Y %H:%M}", GREY),
        ("", None),
        ("Each tab is one backtest run. The top block compares books, the middle "
         "block shows realised P&L by the month a leg was booked, and the bottom "
         "block shows how many legs each exit rule closed and what they made.", None),
        ("", None),
        ("How to read these numbers", BOLD),
        ("Realised is money actually booked. Unrealised is open positions marked "
         "at their last close on the final day of the window — it moves with one "
         "day of market action, so weigh realised more heavily.", None),
        ("Return on capital is measured against each book's own capital, so the "
         "₹2 crore and ₹4 crore rows can be compared directly.", None),
        ("", None),
        ("What the simulation is not", BOLD),
        ("Cash daily bars stand in for stock futures, and entries fill at the "
         "daily close rather than the live 15:15 print, because Angel One serves "
         "only 60 days of futures history per contract. Treat the ranking between "
         "books as the finding and the rupee totals as indicative.", None),
        ("Stops are checked once a day at the close, matching the live book's "
         "cash-close stop rule.", None),
        ("", None),
        ("Source: scripts/backtest_rsi_candle_1y.py", GREY),
    ]
    for r, (text, font) in enumerate(lines, start=1):
        cell = notes.cell(row=r, column=1, value=text)
        if font:
            cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        notes.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        if len(text) > 90:
            notes.row_dimensions[r].height = 30
    notes.column_dimensions["A"].width = 100

    for title, filename, blurb in RUNS:
        path = CACHE / filename
        if not path.exists():
            print(f"  {filename}: not found, skipped")
            continue
        count = write_run(workbook, title, json.loads(path.read_text()), blurb)
        print(f"  {title}: {count} book(s)")

    out = ROOT / "data" / f"RSI_CandlePattern_backtests_{datetime.now():%Y-%m-%d}.xlsx"
    workbook.save(out)
    print(f"\nWrote {out}  ({out.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
